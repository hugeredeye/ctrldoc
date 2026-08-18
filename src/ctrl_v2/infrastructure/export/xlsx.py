from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Font

_ZIP_TIMESTAMP = (1980, 1, 1, 0, 0, 0)
_CREATOR = "CTRL v2"


def _normalize_zip(payload: bytes) -> bytes:
    source = ZipFile(BytesIO(payload), "r")
    output_buffer = BytesIO()
    with source, ZipFile(output_buffer, "w", compression=ZIP_DEFLATED, compresslevel=9) as output:
        for name in sorted(source.namelist()):
            original = source.getinfo(name)
            info = ZipInfo(filename=name, date_time=_ZIP_TIMESTAMP)
            info.compress_type = ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = 0
            output.writestr(info, source.read(name))
    return output_buffer.getvalue()


def _write_header(sheet: object, values: list[str]) -> None:
    sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True)


def export_response_xlsx(snapshot: dict[str, object]) -> bytes:
    workbook = Workbook()
    fixed_time = datetime(2000, 1, 1, tzinfo=UTC)
    workbook.properties.creator = _CREATOR
    workbook.properties.lastModifiedBy = _CREATOR
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False

    summary = workbook.active
    summary.title = "Summary"
    _write_header(summary, ["Metric", "Value"])
    items = list(snapshot["items"])
    summary.append(["Response ID", snapshot["response_id"]])
    summary.append(["RFP ID", snapshot["rfp_id"]])
    summary.append(["Assessment as of", snapshot["assessment_as_of"]])
    summary.append(["Requirements", len(items)])
    for outcome in sorted({str(item["outcome"]) for item in items}):
        summary.append([outcome, sum(1 for item in items if item["outcome"] == outcome)])

    matrix = workbook.create_sheet("Compliance Matrix")
    _write_header(
        matrix,
        [
            "Requirement ID",
            "Order",
            "Requirement",
            "Product",
            "Product Version",
            "Capability",
            "Decision",
            "Outcome",
            "Rationale",
            "Risk",
            "Approved By",
            "Approved At",
        ],
    )
    for item in sorted(items, key=lambda value: (value["source_order"], value["requirement_id"])):
        matrix.append(
            [
                item["requirement_id"],
                item["source_order"],
                item["requirement_text"],
                item["product_name"],
                item["product_version"],
                item["capability_name"],
                item["decision_id"],
                item["outcome"],
                item["rationale"],
                item["risk"],
                item["approved_by"],
                item["approved_at"],
            ]
        )

    evidence_sheet = workbook.create_sheet("Evidence Index")
    _write_header(
        evidence_sheet,
        [
            "Requirement ID",
            "Evidence Span ID",
            "Source Type",
            "Authority",
            "Document Version ID",
            "Document SHA-256",
            "Locator",
            "Exact Quote",
            "Valid From",
            "Valid To",
        ],
    )
    evidence_rows = []
    for item in items:
        for evidence in item["evidence"]:
            evidence_rows.append((item["requirement_id"], evidence))
    for requirement_id, evidence in sorted(
        evidence_rows, key=lambda value: (value[0], value[1]["evidence_span_id"])
    ):
        evidence_sheet.append(
            [
                requirement_id,
                evidence["evidence_span_id"],
                evidence["source_type"],
                evidence["authority_level"],
                evidence["document_version_id"],
                evidence["document_sha256"],
                evidence["locator_canonical"],
                evidence["exact_quote"],
                evidence["valid_from"],
                evidence["valid_to"],
            ]
        )

    metadata = workbook.create_sheet("Metadata")
    _write_header(metadata, ["Field", "Value"])
    for key in (
        "schema_version",
        "response_id",
        "workspace_id",
        "rfp_id",
        "version_no",
        "assessment_as_of",
        "snapshot_hash",
    ):
        metadata.append([key, snapshot[key]])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    raw = BytesIO()
    workbook.save(raw)
    workbook.close()
    return _normalize_zip(raw.getvalue())


class DeterministicXlsxExporter:
    format = "XLSX"
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def export(self, snapshot: dict[str, object]) -> bytes:
        return export_response_xlsx(snapshot)
