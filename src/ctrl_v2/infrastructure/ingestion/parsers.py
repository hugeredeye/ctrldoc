from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document as DocxDocument
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from pypdf import PdfReader


class UnsupportedDocument(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class ParsedBlock:
    ordinal: int
    block_type: str
    text: str
    page_no: int | None
    section_path: tuple[str, ...]
    locator: dict[str, Any]

    @property
    def text_hash(self) -> str:
        return sha256(self.text.encode("utf-8")).hexdigest()


class ParserRegistry:
    contract_version = "1.0"
    parser_build = "stage1-manual"
    supported_media_types = {
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }

    def validate_signature(self, filename: str, media_type: str, content: bytes) -> None:
        suffix = Path(filename).suffix.lower()
        if media_type not in self.supported_media_types:
            raise UnsupportedDocument("Only PDF, DOCX, and XLSX documents are accepted")
        if media_type == "application/pdf":
            if suffix != ".pdf" or not content.startswith(b"%PDF-"):
                raise UnsupportedDocument("The file does not have a valid PDF signature")
            return
        if suffix not in {".docx", ".xlsx"}:
            raise UnsupportedDocument("The extension does not match an OOXML document")
        try:
            with ZipFile(BytesIO(content)) as archive:
                names = set(archive.namelist())
        except BadZipFile as exc:
            raise UnsupportedDocument("The file is not a valid OOXML package") from exc
        required_prefix = "word/" if suffix == ".docx" else "xl/"
        if not any(name.startswith(required_prefix) for name in names):
            raise UnsupportedDocument("The file content does not match its declared type")
        expected = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            if suffix == ".docx"
            else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        if media_type != expected:
            raise UnsupportedDocument("The media type does not match the file extension")

    def parse(self, filename: str, media_type: str, content: bytes) -> list[ParsedBlock]:
        self.validate_signature(filename, media_type, content)
        if media_type == "application/pdf":
            return self._parse_pdf(content)
        if filename.lower().endswith(".docx"):
            return self._parse_docx(content)
        return self._parse_xlsx(content)

    @staticmethod
    def _parse_pdf(content: bytes) -> list[ParsedBlock]:
        reader = PdfReader(BytesIO(content))
        blocks: list[ParsedBlock] = []
        for page_index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                blocks.append(
                    ParsedBlock(
                        ordinal=len(blocks),
                        block_type="PAGE_TEXT",
                        text=text,
                        page_no=page_index,
                        section_path=(),
                        locator={"kind": "PDF", "page": page_index, "bbox": None},
                    )
                )
        return blocks

    @staticmethod
    def _parse_docx(content: bytes) -> list[ParsedBlock]:
        document = DocxDocument(BytesIO(content))
        blocks: list[ParsedBlock] = []
        headings: list[str] = []
        for paragraph_index, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            if style.lower().startswith("heading"):
                try:
                    level = int(style.split()[-1])
                except (ValueError, IndexError):
                    level = 1
                headings = headings[: level - 1] + [text]
            blocks.append(
                ParsedBlock(
                    ordinal=len(blocks),
                    block_type="PARAGRAPH",
                    text=text,
                    page_no=None,
                    section_path=tuple(headings),
                    locator={
                        "kind": "DOCX",
                        "section_path": headings,
                        "paragraph_index": paragraph_index,
                    },
                )
            )
        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    text = cell.text.strip()
                    if text:
                        blocks.append(
                            ParsedBlock(
                                ordinal=len(blocks),
                                block_type="TABLE_CELL",
                                text=text,
                                page_no=None,
                                section_path=tuple(headings),
                                locator={
                                    "kind": "DOCX",
                                    "section_path": headings,
                                    "table_index": table_index,
                                    "row_index": row_index,
                                    "column_index": column_index,
                                },
                            )
                        )
        return blocks

    @staticmethod
    def _parse_xlsx(content: bytes) -> list[ParsedBlock]:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=True)
        blocks: list[ParsedBlock] = []
        try:
            for sheet in workbook.worksheets:
                table_ranges = {
                    table_name: range_boundaries(str(table.ref))
                    for table_name, table in sheet.tables.items()
                }
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        text = str(cell.value).strip()
                        if not text:
                            continue
                        containing_table = next(
                            (
                                name
                                for name, bounds in table_ranges.items()
                                if bounds[0] <= cell.column <= bounds[2]
                                and bounds[1] <= cell.row <= bounds[3]
                            ),
                            None,
                        )
                        blocks.append(
                            ParsedBlock(
                                ordinal=len(blocks),
                                block_type="CELL",
                                text=text,
                                page_no=None,
                                section_path=(sheet.title,),
                                locator={
                                    "kind": "XLSX",
                                    "sheet": sheet.title,
                                    "cell_range": cell.coordinate,
                                    "table_name": containing_table,
                                },
                            )
                        )
        finally:
            workbook.close()
        return blocks
