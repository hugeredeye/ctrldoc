from __future__ import annotations

from io import BytesIO

from docx import Document
from openpyxl import Workbook

from ctrl_v2.infrastructure.ingestion import ParserRegistry


def _minimal_pdf(text: str) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        (
            f"<< /Length {len(text) + 33} >>\nstream\n"
            f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET\nendstream"
        ).encode("ascii"),
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(payload)


def test_pdf_parser_preserves_page_provenance():
    blocks = ParserRegistry().parse("spec.pdf", "application/pdf", _minimal_pdf("SAML support"))

    assert blocks[0].text == "SAML support"
    assert blocks[0].locator == {"kind": "PDF", "page": 1, "bbox": None}


def test_docx_parser_preserves_section_and_table_provenance():
    document = Document()
    document.add_heading("Authentication", level=1)
    document.add_paragraph("SAML support")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Certified"
    output = BytesIO()
    document.save(output)

    blocks = ParserRegistry().parse(
        "spec.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        output.getvalue(),
    )

    paragraph = next(block for block in blocks if block.text == "SAML support")
    table_cell = next(block for block in blocks if block.text == "Certified")
    assert paragraph.section_path == ("Authentication",)
    assert table_cell.locator["table_index"] == 0
    assert table_cell.locator["row_index"] == 0
    assert table_cell.locator["column_index"] == 0


def test_xlsx_parser_preserves_sheet_and_cell_provenance():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Capabilities"
    sheet["C7"] = "SAML support"
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    blocks = ParserRegistry().parse(
        "spec.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        output.getvalue(),
    )

    assert blocks[0].locator["sheet"] == "Capabilities"
    assert blocks[0].locator["cell_range"] == "C7"
